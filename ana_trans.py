#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Synchronisation entre paroles (transcription) et fixations du regard.

Usage :
    python synchro_paroles_regard.py anv_psy.csv anv_sujet.csv transcription.txt

Options :
    --entretien NOM      nom de l'entretien utilisé dans le nom des fichiers créés
                         (par défaut : nom du fichier de transcription, sans
                         le suffixe "_eye_tracker")
    --sortie DOSSIER     dossier de sortie (par défaut : celui de la transcription)
    --decalage SECONDES  décalage ajouté aux débuts/fins de parole (défaut : 0)
    --sans-silences      ne garder que les instants où quelqu'un parle
    --locuteurs ...      locuteurs à prendre en compte (défaut : psy sujet br)
    --sep SEP            séparateur de la transcription (défaut : détecté ; "tab")
    --graphes            enregistre aussi le graphe de répartition des locuteurs

Entrées :
    - anv_psy.csv, anv_sujet.csv : produits par les scripts d'analyse du regard
      (colonnes Capture.Time, region, saccade et, si disponible, type_absence)
    - transcription : fichier SANS en-tête dont les colonnes 1, 4, 6, 8 et 9 sont
      le locuteur, le début, la fin, la durée (en secondes) et la parole

Sorties (dans le même dossier que la transcription) :
    - sync_FP_TAbs_<entretien>.csv
    - sync_FP_TAbs_<entretien>.xlsx   (mis en forme avec couleurs)

Dépendances : pandas, numpy, matplotlib, openpyxl
"""

import argparse
import csv
import os
import re
import sys
from collections import Counter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------------------
# Paramètres par défaut
# ------------------------------------------------------------------------------
LOCUTEURS_DEFAUT = ["psy", "sujet", "br"]   # retirer "br" pour ignorer le bruit

TITRE_TEMPS = "Répartition du temps de parole de chaque locuteurs"
TITRE_NB = "Nombre de prise de parole totale de chaque locuteurs"

PALETTE_REGARDEUR = {"psy": "9DC3E6", "sujet": "E6A8C8"}
PALETTE_REGION = {
    # codes numériques
    "1": "A9E5A0", "2": "F4B6D9", "3": "FFF2A8", "4": "F4A6A6",
    "5": "A6C4F0", "254": "D9D9D9", "255": "D9D9D9",
    # étiquettes
    "Yeux": "A9E5A0", "Bouche_Nez": "F4B6D9", "Nez_et_bouche": "F4B6D9",
    "Visage": "FFF2A8", "Corps": "F4A6A6", "Autre": "A6C4F0",
    "Frame_invalide": "D9D9D9", "Absence_valeurs": "D9D9D9", "Absence_valeur": "D9D9D9",
}
COULEUR_LOC = {"psy": "DDEBF7", "sujet": "F8E0EC", "br": "EDEDED"}


# ------------------------------------------------------------------------------
# Lecture des fichiers
# ------------------------------------------------------------------------------
def detecter_separateur(chemin, encodage):
    """Détecte le séparateur (tabulation, ; ou ,) sur les premières lignes."""
    with open(chemin, encoding=encodage, errors="replace", newline="") as f:
        echantillon = [l for l in (f.readline() for _ in range(30)) if l.strip()]
    meilleur, meilleur_score = ",", (-1.0, -1)
    for cand in ("\t", ";", ","):
        longueurs = [len(r) for r in csv.reader(echantillon, delimiter=cand)]
        if not longueurs:
            continue
        modale, freq = Counter(longueurs).most_common(1)[0]
        if modale < 2:
            continue
        if freq == len(longueurs):      # séparateur parfaitement régulier
            return cand
        score = (freq / len(longueurs), modale)
        if score > meilleur_score:
            meilleur, meilleur_score = cand, score
    return meilleur


def lire_table(chemin, header, sep=None):
    """Lit un fichier texte tabulaire (tout en texte ; vide / "NA" -> manquant)."""
    for enc in ("utf-8-sig", "latin-1"):
        try:
            s = sep or detecter_separateur(chemin, enc)
            options = {"quoting": csv.QUOTE_NONE} if s == "\t" else {}
            return pd.read_csv(chemin, sep=s, header=header, dtype=str, encoding=enc,
                               keep_default_na=False, na_values=["", "NA"], **options)
        except UnicodeDecodeError:
            continue
    sys.exit(f"Impossible de lire {chemin} (encodage non reconnu).")


def en_numerique(serie):
    """Convertit en nombres (accepte la virgule décimale)."""
    return pd.to_numeric(serie.str.replace(",", ".", regex=False), errors="coerce")


def lire_transcription(chemin, sep=None):
    """Retourne boc_trans : locuteur, debut, fin, duree, parole (trié par début)."""
    X = lire_table(chemin, header=None, sep=sep)
    if X.shape[1] < 9:
        sys.exit(f"La transcription doit avoir au moins 9 colonnes "
                 f"(trouvées : {X.shape[1]}).")
    t = X.iloc[:, [0, 3, 5, 7, 8]].copy()          # V1, V4, V6, V8, V9
    t.columns = ["locuteur", "debut", "fin", "duree", "parole"]
    for c in ["debut", "fin", "duree"]:
        t[c] = en_numerique(t[c])
    t["locuteur"] = t["locuteur"].str.strip()
    t = t.dropna(subset=["debut"])                  # écarte aussi une éventuelle ligne d'en-tête
    return t.sort_values("debut", kind="stable").reset_index(drop=True)


def lire_regard(chemin, nom):
    """Lit un fichier anv_psy / anv_sujet."""
    df = lire_table(chemin, header=0)
    df.columns = [str(c).strip().replace(" ", ".") for c in df.columns]
    manquantes = [c for c in ("Capture.Time", "region", "saccade") if c not in df.columns]
    if manquantes:
        sys.exit(f"{os.path.basename(chemin)} : colonnes introuvables {manquantes}\n"
                 f"Colonnes disponibles : {list(df.columns)}")
    df["Capture.Time"] = en_numerique(df["Capture.Time"])
    if "type_absence" not in df.columns:
        print(f"ATTENTION : pas de colonne 'type_absence' dans {os.path.basename(chemin)} "
              f"-> colonne type_absence_{nom} laissée vide.")
        df["type_absence"] = np.nan
    return df


# ------------------------------------------------------------------------------
# Statistiques et graphe des locuteurs
# ------------------------------------------------------------------------------
def stats_locuteurs(trans):
    t = trans.dropna(subset=["locuteur"])
    recap = (t.groupby("locuteur")
             .agg(nb_loc=("duree", "size"),
                  temps_total=("duree", "sum"),
                  duree_moyenne=("duree", "mean"))
             .reset_index())
    recap["pct_temps"] = recap["temps_total"] / recap["temps_total"].sum() * 100
    recap["pct_loc"] = recap["nb_loc"] / recap["nb_loc"].sum() * 100
    return recap.sort_values("pct_temps", ascending=False).reset_index(drop=True)


def barh_etiquettes(ax, labels, valeurs, couleur, etiquettes, xlim=None):
    ax.barh(labels, valeurs, color=couleur, height=0.6)
    vmax = max(valeurs) if len(valeurs) else 1
    for i, txt in enumerate(etiquettes):
        ax.text(valeurs[i] + vmax * 0.01, i, txt, va="center", fontsize=11)
    ax.set_xlim(*(xlim if xlim else (0, vmax * 1.15)))
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.xaxis.grid(True, color="#e5e5e5")
    ax.set_axisbelow(True)


def graphe_locuteurs(recap, chemin_png):
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 5))

    d1 = recap.sort_values("pct_temps")
    barh_etiquettes(ax1, d1["locuteur"].tolist(), d1["pct_temps"].tolist(), "#2a78d6",
                    [f"{v:.1f}%" for v in d1["pct_temps"]], xlim=(0, 100))
    ax1.set_title(TITRE_TEMPS, loc="left", fontsize=13)
    ax1.set_xlabel("% du temps total")

    d2 = recap.sort_values("nb_loc")
    barh_etiquettes(ax2, d2["locuteur"].tolist(), d2["nb_loc"].tolist(), "#1baf7a",
                    [f"{v:.1f}%" for v in d2["pct_loc"]])
    ax2.set_title(TITRE_NB, loc="left", fontsize=13)
    ax2.set_xlabel("Nombre de prises de paroles")

    fig.tight_layout()
    fig.savefig(chemin_png, dpi=300)
    plt.close(fig)


# ------------------------------------------------------------------------------
# Synchronisation paroles / regard
# ------------------------------------------------------------------------------
def differe(x):
    """True à chaque changement de valeur (NA compris) ; True sur la 1re ligne."""
    x = pd.Series(list(x))
    p = x.shift(1)
    r = (x.isna() != p.isna()) | (x.notna() & p.notna() & (x != p))
    r = np.array(r, dtype=bool)  # copie modifiable
    if len(r):
        r[0] = True
    return r


def construire_frise(regard, trans, regardeur, locuteurs, decalage, inclure_silences):
    ids = [f"id_{l}" for l in locuteurs]

    # 1) Intervalles de parole
    iv = trans.copy()
    iv["id"] = np.arange(1, len(iv) + 1)
    iv["t_debut"] = iv["debut"] + decalage
    iv["t_fin"] = iv["fin"] + decalage
    iv = iv[iv["locuteur"].isin(locuteurs) & iv["t_debut"].notna() & iv["t_fin"].notna()]
    if iv.empty:
        sys.exit(f"Aucune parole trouvée pour les locuteurs {locuteurs} dans la transcription.")

    # 2) Échantillons de regard : chacun couvre [t, t_suivant[
    g = pd.DataFrame({"t": regard["Capture.Time"] / 1e6,
                      "region": regard["region"],
                      "saccade": regard["saccade"]}).dropna(subset=["t"])
    g = (g.sort_values("t", kind="stable")
          .drop_duplicates("t", keep="first")
          .reset_index(drop=True))
    t = g["t"].to_numpy()
    dt = float(np.median(np.diff(t))) if len(t) > 1 else 0.0
    g["t_next"] = np.append(t[1:], t[-1] + dt)

    if inclure_silences:
        g = g[(g["t"] >= iv["t_debut"].min()) & (g["t"] <= iv["t_fin"].max())]
        g = g.reset_index(drop=True)
        t = g["t"].to_numpy()

    # 3) Ce que dit chaque locuteur à chaque instant (colonnes id_ et parole_)
    for l in locuteurs:
        ivl = iv[iv["locuteur"] == l].sort_values("id")
        id_arr = np.full(len(t), np.nan)
        parole_arr = np.full(len(t), None, dtype=object)
        for row in ivl.itertuples(index=False):
            lo = np.searchsorted(t, row.t_debut, side="left")   # t >= t_debut
            hi = np.searchsorted(t, row.t_fin, side="left")     # t <  t_fin
            if hi > lo:
                # l'énoncé d'id le plus petit l'emporte en cas de chevauchement
                libres = np.flatnonzero(np.isnan(id_arr[lo:hi])) + lo
                id_arr[libres] = row.id
                parole_arr[libres] = row.parole
        g[f"id_{l}"] = id_arr
        g[f"parole_{l}"] = parole_arr

    if not inclure_silences:
        g = g[g[ids].notna().any(axis=1)].reset_index(drop=True)

    if g.empty:
        sys.exit(f"Aucun échantillon de regard ({regardeur}) ne coïncide avec la transcription.")

    # 4) Nouveau bloc dès que la région OU l'un des énoncés change
    changement = np.zeros(len(g), dtype=bool)
    for k in ["region"] + ids:
        changement |= differe(g[k])
    g["bloc"] = np.cumsum(changement)

    agg = {"debut_t": ("t", "min"), "fin_t": ("t_next", "max"),
           "region": ("region", "first"), "n_mesures": ("t", "size")}
    for k in ids:
        agg[k] = (k, "first")
    for l in locuteurs:
        agg[f"parole_{l}"] = (f"parole_{l}", "first")

    groupes = g.groupby("bloc", sort=True)
    blocs = groupes.agg(**agg)
    blocs["saccade"] = groupes["saccade"].agg(lambda s: "/".join(dict.fromkeys(s.dropna())))
    blocs = blocs.reset_index(drop=True)

    # 5) Anti-répétition : le texte n'apparaît que sur la 1re ligne de l'énoncé
    for l in locuteurs:
        idc, pc = f"id_{l}", f"parole_{l}"
        premiere = differe(blocs[idc]) & blocs[idc].notna().to_numpy()
        blocs[pc] = blocs[pc].where(premiere, np.nan)

    # Qui parle pendant le bloc (même quand le texte n'est pas répété)
    present = np.column_stack([blocs[k].notna().to_numpy() for k in ids])
    blocs["qui_parle"] = ["+".join(l for l, m in zip(locuteurs, ligne) if m) or "silence"
                          for ligne in present]

    blocs["regardeur"] = regardeur
    blocs["duree"] = (blocs["fin_t"] - blocs["debut_t"]).round(3)
    blocs["debut_t"] = blocs["debut_t"].round(3)
    blocs["fin_t"] = blocs["fin_t"].round(3)
    blocs.loc[blocs["saccade"] == "", "saccade"] = np.nan

    colonnes = (["regardeur", "debut_t", "fin_t", "duree", "region", "saccade", "qui_parle"]
                + [f"parole_{l}" for l in locuteurs] + ids + ["n_mesures"])
    return blocs[colonnes]


def types_absence_par_ligne(res, regard):
    """Types d'absence de ce regardeur dont l'échantillon tombe dans [debut_t, fin_t[."""
    g = pd.DataFrame({"t": (regard["Capture.Time"] / 1e6).round(3),
                      "type": regard["type_absence"]}).dropna()
    g = g.sort_values("t", kind="stable")
    t = g["t"].to_numpy()
    types = g["type"].to_numpy(dtype=object)
    lo = np.searchsorted(t, res["debut_t"].to_numpy(), side="left")   # debut_t <= t
    hi = np.searchsorted(t, res["fin_t"].to_numpy(), side="left")     # t < fin_t
    sortie = [("/".join(dict.fromkeys(types[a:b])) if b > a else np.nan)
              for a, b in zip(lo, hi)]
    return pd.Series(sortie, index=res.index, dtype=object)


# ------------------------------------------------------------------------------
# Export Excel (couleurs)
# ------------------------------------------------------------------------------
def exporter_excel(res, chemin, locuteurs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    noms = list(res.columns)
    nc, nr = len(noms), len(res)
    idx = {n: j for j, n in enumerate(noms, start=1)}

    # Données (les valeurs manquantes restent des cellules vides)
    ws.append(noms)
    for ligne in res.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)
                   for v in ligne])

    # Styles
    gris = Side(style="thin", color="BFBFBF")
    noir = Side(style="thin", color="000000")
    bord = Border(left=gris, right=gris, top=gris, bottom=gris)
    bord_ent = Border(left=noir, right=noir, top=noir, bottom=noir)
    rempl = lambda hexa: PatternFill(fill_type="solid", fgColor=hexa)
    al_haut = Alignment(vertical="top")
    al_retour = Alignment(wrap_text=True, vertical="top")
    al_centre = Alignment(horizontal="center", vertical="top")
    al_droite = Alignment(horizontal="right", vertical="top")
    gras = Font(bold=True)

    # En-tête
    for j in range(1, nc + 1):
        c = ws.cell(row=1, column=j)
        c.fill = rempl("1F3864")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = bord_ent

    cols_parole = [idx[f"parole_{l}"] for l in locuteurs]
    cols_num = [idx[n] for n in ("debut_t", "fin_t", "duree")]
    v_regardeur = res["regardeur"].tolist()
    v_region = res["region"].tolist()
    v_saccade = res["saccade"].tolist()
    v_parole = {l: res[f"parole_{l}"].tolist() for l in locuteurs}

    for i in range(nr):
        ligne = i + 2
        for j in range(1, nc + 1):
            c = ws.cell(row=ligne, column=j)
            c.border = bord
            c.alignment = al_haut
        for j in cols_parole:
            ws.cell(row=ligne, column=j).alignment = al_retour
        for j in cols_num:
            c = ws.cell(row=ligne, column=j)
            c.number_format = "0.000"
            c.alignment = al_droite

        if v_regardeur[i] in PALETTE_REGARDEUR:
            c = ws.cell(row=ligne, column=idx["regardeur"])
            c.fill, c.font, c.alignment = rempl(PALETTE_REGARDEUR[v_regardeur[i]]), gras, al_centre
        if v_region[i] in PALETTE_REGION:
            c = ws.cell(row=ligne, column=idx["region"])
            c.fill, c.alignment = rempl(PALETTE_REGION[v_region[i]]), al_centre
        for l in locuteurs:
            if pd.notna(v_parole[l][i]):
                c = ws.cell(row=ligne, column=idx[f"parole_{l}"])
                c.fill = rempl(COULEUR_LOC.get(l, "EDEDED"))
                c.alignment = al_retour
        if pd.notna(v_saccade[i]):
            c = ws.cell(row=ligne, column=idx["saccade"])
            c.fill, c.alignment = rempl("FFD966"), al_centre

    # Largeurs, hauteur d'en-tête, volet figé, filtre
    for j in range(1, nc + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12
    for j in cols_parole:
        ws.column_dimensions[get_column_letter(j)].width = 45
    ws.column_dimensions[get_column_letter(idx["qui_parle"])].width = 14
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(nc)}{nr + 1}"

    wb.save(chemin)


# ------------------------------------------------------------------------------
# Programme principal
# ------------------------------------------------------------------------------
def nom_entretien(chemin_transcription, explicite=None):
    if explicite:
        nom = explicite
    else:
        nom = os.path.splitext(os.path.basename(chemin_transcription))[0]
        nom = re.sub(r"[_\-. ]*eye[_\-. ]*tracker$", "", nom, flags=re.I)
    return re.sub(r"[^\w.\-]+", "_", nom).strip("_") or "entretien"


def main():
    p = argparse.ArgumentParser(
        description="Synchronisation entre paroles et fixations du regard.")
    p.add_argument("anv_psy", help="Fichier anv_psy.csv")
    p.add_argument("anv_sujet", help="Fichier anv_sujet.csv")
    p.add_argument("transcription", help="Fichier de transcription de l'entretien")
    p.add_argument("--entretien", help="Nom de l'entretien (utilisé dans le nom des sorties)")
    p.add_argument("--sortie", help="Dossier de sortie (défaut : dossier de la transcription)")
    p.add_argument("--decalage", type=float, default=0.0,
                   help="Secondes ajoutées aux débuts/fins de parole (défaut : 0)")
    p.add_argument("--sans-silences", action="store_true",
                   help="Ne garder que les instants où quelqu'un parle")
    p.add_argument("--locuteurs", nargs="+", default=LOCUTEURS_DEFAUT,
                   help="Locuteurs à prendre en compte (défaut : psy sujet br)")
    p.add_argument("--sep", help='Séparateur de la transcription (défaut : détecté ; "tab")')
    p.add_argument("--graphes", action="store_true",
                   help="Enregistre aussi le graphe de répartition des locuteurs")
    args = p.parse_args()

    locuteurs = args.locuteurs
    inclure_silences = not args.sans_silences
    sep = "\t" if args.sep in ("tab", "\\t") else args.sep

    chemin_trans = os.path.abspath(args.transcription)
    entretien = nom_entretien(chemin_trans, args.entretien)
    out_dir = os.path.abspath(args.sortie) if args.sortie else os.path.dirname(chemin_trans)
    os.makedirs(out_dir, exist_ok=True)
    nname = f"sync_FP_TAbs_{entretien}.csv"

    # Entrées
    boc_trans = lire_transcription(chemin_trans, sep)
    XP = lire_regard(args.anv_psy, "psy")
    XS = lire_regard(args.anv_sujet, "sujet")
    regardeurs = {"psy": XP, "sujet": XS}

    trouves = set(boc_trans["locuteur"].dropna())
    for l in locuteurs:
        if l not in trouves:
            print(f"ATTENTION : le locuteur '{l}' est absent de la transcription "
                  f"(locuteurs trouvés : {sorted(trouves)}).")

    # Statistiques sur les locuteurs
    recap_loc = stats_locuteurs(boc_trans)
    print(recap_loc.to_string(index=False))
    if args.graphes:
        graphe_locuteurs(recap_loc, os.path.join(out_dir, f"comparaison_locuteurs_{entretien}.png"))

    # Construction
    res = pd.concat([construire_frise(regard, boc_trans, nom, locuteurs,
                                      args.decalage, inclure_silences)
                     for nom, regard in regardeurs.items()], ignore_index=True)
    res = res.sort_values(["regardeur", "debut_t"], kind="stable").reset_index(drop=True)

    # Type d'absence du psy et du sujet sur chaque ligne du résultat
    res["type_absence_psy"] = types_absence_par_ligne(res, XP)
    res["type_absence_sujet"] = types_absence_par_ligne(res, XS)
    colonnes = list(res.columns)
    for c in ("type_absence_psy", "type_absence_sujet"):
        colonnes.remove(c)
    pos = colonnes.index("saccade") + 1
    colonnes[pos:pos] = ["type_absence_psy", "type_absence_sujet"]
    res = res[colonnes]
    for l in locuteurs:
        res[f"id_{l}"] = res[f"id_{l}"].astype("Int64")

    # Contrôle de continuité : chaque début doit égaler la fin de la ligne précédente
    print()
    for r in res["regardeur"].unique():
        d = res[res["regardeur"] == r]
        ecart = d["debut_t"].to_numpy()[1:] - d["fin_t"].to_numpy()[:-1]
        print(f"{r:<6} lignes={len(d)} | trous/chevauchements (|ecart|>0.002 s) = "
              f"{int((np.abs(ecart) > 0.002).sum())}")

    # Export Excel (couleurs) + CSV
    chemin_csv = os.path.join(out_dir, nname)
    chemin_xlsx = os.path.join(out_dir, re.sub(r"\.csv$", ".xlsx", nname))
    exporter_excel(res, chemin_xlsx, locuteurs)
    res.to_csv(chemin_csv, index=False, na_rep="", encoding="utf-8")

    print(f"\nFichiers enregistrés dans : {out_dir}")
    print(f"  - {os.path.basename(chemin_csv)}")
    print(f"  - {os.path.basename(chemin_xlsx)}")
    if args.graphes:
        print(f"  - comparaison_locuteurs_{entretien}.png")


if __name__ == "__main__":
    main()
